"""Genera un reporte de stock ficticio con el mismo formato que espera el archivo.

Uso:  python tools/generar_demo.py
Salida: demo/Aurora/reporte_demo.xlsx

Los datos son inventados: marcas, SKU, seriales y bodegas no corresponden a
ninguna empresa real. Sirven para probar el flujo completo sin datos reales.
"""

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "demo" / "Aurora" / "reporte_demo.xlsx"

# El archivo de inventario exige estas columnas; el resto es relleno realista.
COLUMNAS = [
    "Serial", "Código", "Descripción", "Condición", "Estado", "Grado",
    "Picking", "Nota de Venta", "Bodega", "Ubicación", "Categoría",
    "Sucursal", "Tipo de Bodega", "Fecha de Ingreso a Bodega", "Observación",
]

SUCURSAL_PRINCIPAL = "Casa Matriz"
SUCURSALES = [(SUCURSAL_PRINCIPAL, 420), ("Sucursal Norte", 150)]

PRODUCTOS = [
    # (SKU, descripción, categoría, prefijo de serial, largo del serial)
    ("AX-NB14-512", 'Notebook Aurex Pro 14", Ryzen 7, 16GB/512GB SSD', "Notebooks", "AXN", 10),
    ("AX-NB15-256", 'Notebook Aurex Lite 15", Core i5, 8GB/256GB SSD', "Notebooks", "AXL", 10),
    ("KB-AIO24-01", 'All-in-One Kobalt 24", Core i7, 16GB/1TB SSD', "Desktops", "KBA", 12),
    ("KB-MT500", "Desktop Kobalt Tower MT500, Core i5, 8GB/512GB SSD", "Desktops", "KBM", 12),
    ("LM-MON27", 'Monitor Lumen 27" QHD 75Hz', "Monitores", "LMM", 11),
    ("LM-MON24", 'Monitor Lumen 24" FHD', "Monitores", "LMD", 11),
    ("VT-IMP450", "Impresora Vantis LaserJet 450dn", "Impresoras", "VTI", 10),
    ("VT-MFP620", "Multifuncional Vantis 620dnw", "Impresoras", "VTM", 10),
    ("ND-DOCK130", "Docking Station Nordika USB-C 130W", "Accesorios", "NDD", 13),
    ("ND-TEC-INA", "Teclado y mouse inalámbrico Nordika", "Accesorios", "NDT", 13),
    ("AX-BAT-4C", "Batería 4 celdas - Notebook Aurex Pro", "Partes y Piezas", "", 14),
    ("AX-LCD-14F", 'Panel LCD 14" FHD - Notebook Aurex Pro', "Partes y Piezas", "", 14),
    ("KB-RAM-8G", "Memoria RAM DDR4 8GB 3200MHz", "Partes y Piezas", "", 14),
    ("KB-SSD-512", "Disco SSD NVMe 512GB", "Partes y Piezas", "", 14),
    ("VT-TON-45A", "Tóner Vantis 45A negro", "Suministros", "VTT", 10),
    ("ND-SILLA-ER", "Silla ergonómica Nordika malla", "Muebles", "NDS", 9),
    ("ND-ESC-140", "Escritorio Nordika 140x70 cm", "Muebles", "NDE", 9),
    ("LM-PROY-HD", "Proyector Lumen HD 3000 lúmenes", "Audio y Video", "LMP", 11),
]

BODEGAS = [
    ("Almacén Principal", "Almacen", 0.34),
    ("Partes y Piezas", "Almacen", 0.22),
    ("Activo Fijo", "Almacen", 0.16),
    ("Ventas", "Almacen", 0.10),
    ("Laboratorio", "Servicio", 0.08),
    ("Exhibición", "Almacen", 0.05),
    ("Desarme", "Almacen", 0.05),
]

CONDICIONES = ["Usado"] * 6 + ["Nuevo"] * 3 + ["Reacondicionado", "Defectuoso"]
ESTADOS = ["Operativo", "Disponible", "Asignado", "Sin Uso", "Por Revisar", "Partes y Piezas"]
GRADOS = ["No"] * 7 + ["A", "B", "C"]
OBSERVACIONES = [
    "", "", "", "", "",
    "Requiere limpieza",
    "Equipo revisado, operativo",
    "Falta cable de poder",
    "Caja abierta",
    "Pendiente de diagnóstico",
]

LETRAS = "ABCDEFGHJKLMNPQRSTUVWXYZ"
DIGITOS = "0123456789"


def serial(rnd, prefijo, largo):
    """Genera seriales de formatos variados, incluidos numéricos con ceros."""
    if not prefijo:
        # Seriales solo numéricos: ponen a prueba el manejo de ceros iniciales
        return str(rnd.randint(1, 9_999_999)).zfill(largo)
    cuerpo = "".join(rnd.choice(LETRAS + DIGITOS) for _ in range(largo - len(prefijo)))
    return prefijo + cuerpo


def ubicacion(rnd):
    return f"P{rnd.randint(1, 2)}-E{rnd.randint(1, 4)}-R{rnd.randint(1, 4)}-N{rnd.randint(1, 4)}"


def generar():
    rnd = random.Random(20260914)
    filas = []
    usados = set()
    nv = 12500
    picking = 30000

    for sucursal, cantidad in SUCURSALES:
        for _ in range(cantidad):
            sku, desc, categoria, prefijo, largo = rnd.choice(PRODUCTOS)
            s = serial(rnd, prefijo, largo)
            while s in usados:
                s = serial(rnd, prefijo, largo)
            usados.add(s)

            bodega, tipo, _peso = rnd.choices(BODEGAS, weights=[b[2] for b in BODEGAS])[0]
            # Unos pocos equipos van de salida: tienen Nota de Venta y Picking
            de_salida = sucursal == SUCURSAL_PRINCIPAL and rnd.random() < 0.02
            if de_salida:
                nv += 1
                picking += 1

            filas.append([
                s,
                sku,
                desc,
                rnd.choice(CONDICIONES),
                rnd.choice(ESTADOS),
                rnd.choice(GRADOS),
                str(picking) if de_salida else "No",
                str(nv) if de_salida else "No",
                bodega,
                ubicacion(rnd),
                categoria,
                sucursal,
                tipo,
                f"{rnd.randint(1, 28):02d}-{rnd.randint(1, 12):02d}-202{rnd.randint(4, 6)}",
                rnd.choice(OBSERVACIONES),
            ])

    rnd.shuffle(filas)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(COLUMNAS)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append(fila)
    # El serial va como texto: si Excel lo tomara como número, perdería los ceros
    for celda in ws["A"][1:]:
        celda.number_format = "@"
    ws.freeze_panes = "A2"
    for col, ancho in zip("ABCDEFGHIJKLMNO", [18, 14, 52, 16, 14, 8, 10, 14, 22, 18, 18, 16, 14, 18, 30]):
        ws.column_dimensions[col].width = ancho

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)

    principal = [f for f in filas if f[11] == SUCURSAL_PRINCIPAL]
    salida_pendiente = [f for f in principal if f[7] != "No"]
    print(f"Generado: {SALIDA}")
    print(f"  Filas totales: {len(filas)}")
    print(f"  {SUCURSAL_PRINCIPAL}: {len(principal)}  (de salida: {len(salida_pendiente)})")
    print(f"  Esperados a contar: {len(principal) - len(salida_pendiente)}")


if __name__ == "__main__":
    generar()
